from typing import List, Dict, Any, Tuple
from pathlib import Path
import openpyxl

from schemas.common import DocumentSource, DocumentStatus, DocumentType
from schemas.document import DocumentChunk
from schemas.technical import TechnicalRecord
from src.ingestion.metadata import normalize_revision, normalize_status, normalize_unit, resolve_equipment


def normalize_plot_plan_excel(
    file_path: Path
) -> Tuple[List[TechnicalRecord], List[DocumentChunk]]:
    """
    Parses Plot Plan / Layout Excel into:
    1. TechnicalRecord[] (Grid Reference, Elevation, Coordinates)
    2. DocumentChunk[] (narrative layout and nearby equipment)
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    meta: Dict[str, str] = {}
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

    raw_tag = meta.get("equipment_tag")
    res = resolve_equipment(raw_tag)
    tag = res["equipment_tag"] if res.get("status") == "resolved" else (raw_tag or "Asset")
    doc_id = meta.get("document_no") or meta.get("document_id") or f"PP-{tag}"

    # Never infer revision if missing
    rev = normalize_revision(meta.get("revision"))
    status = normalize_status(meta.get("revision_status")) or DocumentStatus.ISSUED_FOR_CONSTRUCTION
    status_str = status.value if hasattr(status, "value") else str(status)
    source_file = meta.get("source_file") or file_path.name.replace(".xlsx", ".pdf")

    source = DocumentSource(
        file_name=source_file,
        page=1,
        sheet="Location Data",
        revision=rev,
        status=status
    )

    technical_records: List[TechnicalRecord] = []
    loc_lines: List[str] = []
    nearby_lines: List[str] = []

    if "Location Data" in wb.sheetnames:
        ws = wb["Location Data"]
        for r in range(3, 15):
            param = ws.cell(r, 1).value
            val = ws.cell(r, 2).value
            unit = ws.cell(r, 3).value
            if param and val:
                p_str = str(param).strip()
                v_str = str(val).strip()
                u_str = normalize_unit(str(unit).strip()) if unit and str(unit).strip() not in ("—", "-") else None

                # Create TechnicalRecord for coordinates & elevation
                canon_f = None
                if "grid" in p_str.lower():
                    canon_f = "grid_reference"
                elif "elevation" in p_str.lower():
                    canon_f = "elevation"

                tr = TechnicalRecord(
                    record_id=f"TR-{tag}-{p_str.upper().replace(' ', '_')}",
                    record_type="coordinate",
                    equipment_tag=tag,
                    document_id=doc_id,
                    document_type=DocumentType.PLOT_PLAN,
                    category="Plant Location",
                    parameter=p_str,
                    canonical_field=canon_f,
                    source_label=p_str,
                    value=v_str,
                    unit=u_str,
                    source=source
                )
                technical_records.append(tr)
                u_info = f" {u_str}" if u_str else ""
                loc_lines.append(f"{p_str}: {v_str}{u_info}")

        # Nearby equipment
        for r in range(17, ws.max_row + 1):
            etag = ws.cell(r, 1).value
            ename = ws.cell(r, 2).value
            area_code = ws.cell(r, 3).value
            area_name = ws.cell(r, 4).value
            el = ws.cell(r, 5).value
            if etag:
                name_str = f" ({ename})" if ename and str(ename).strip() not in ("-", "—", "") else ""
                nearby_lines.append(f"- {etag}{name_str}: Area {area_code} ({area_name}), Elevation {el}")

    content_lines = [
        f"PLOT PLAN & EQUIPMENT LOCATION DATA: {tag} ({meta.get('equipment_name', 'Equipment')}).",
        f"Document No: {doc_id}, Revision: {rev or 'N/A'}, Status: {status_str}.",
        f"Area: {meta.get('area', '')}, Functional Location: {meta.get('functional_location', '')}.",
        "LOCATION SPECIFICATIONS & COORDINATES:",
        "; ".join(loc_lines),
        "NEARBY PLANT EQUIPMENT & ELEVATIONS:",
        "\n".join(nearby_lines)
    ]
    chunk = DocumentChunk(
        chunk_id=f"{doc_id}-P01-C01",
        document_id=doc_id,
        document_type=DocumentType.PLOT_PLAN,
        title=f"Plot Plan & Location Data - {tag}",
        equipment_tag=tag,
        content="\n".join(content_lines),
        source=source,
        metadata={"total_nearby_equipment": len(nearby_lines)}
    )

    return technical_records, [chunk]
