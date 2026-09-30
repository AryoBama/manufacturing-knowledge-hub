from typing import List, Dict, Any
from pathlib import Path
import openpyxl

from schemas.common import DocumentSource, DocumentStatus, DocumentType
from schemas.document import DocumentChunk
from src.ingestion.metadata import normalize_revision, normalize_status, resolve_equipment


def normalize_opl_sheet_dict(
    sheet_data: Dict[str, Any],
    source_metadata: DocumentSource
) -> DocumentChunk:
    doc_id = sheet_data.get("document_id") or sheet_data.get("opl_no", "OPL-GEN")
    title = sheet_data.get("title") or sheet_data.get("opl_title", "One Point Lesson")
    tag = sheet_data.get("equipment_tag")
    if tag:
        res = resolve_equipment(tag)
        if res.get("status") == "resolved":
            tag = res["equipment_tag"]

    disc = sheet_data.get("discipline", "Mechanical / Operations")
    area = sheet_data.get("area_unit") or sheet_data.get("area", "")
    interlock_ref = sheet_data.get("related_interlock", "")
    pid_ref = sheet_data.get("pid_reference", "")

    sections = sheet_data.get("sections", {})
    purpose = sections.get("purpose", "")
    precautions = sections.get("safety_precautions", [])
    tools = sections.get("tools", [])
    steps = sections.get("steps", [])
    troubleshooting = sections.get("troubleshooting", [])
    learnings = sections.get("learnings", [])

    content_lines = [
        f"ONE POINT LESSON (OPL): {doc_id} - {title}.",
        f"Equipment: {tag or 'Plant Equipment'}, Discipline: {disc}, Area: {area}.",
        f"References: Interlock: {interlock_ref}, P&ID: {pid_ref}.",
        f"PURPOSE & OBJECTIVE:\n{purpose}",
        "SAFETY PRECAUTIONS:\n" + "\n".join(f"- {p}" for p in precautions),
        "TOOLS & MATERIALS:\n" + "\n".join(f"- {t}" for t in tools),
        "STEP-BY-STEP PROCEDURE:\n" + "\n".join(f"- {s}" for s in steps),
        "COMMON PROBLEMS & TROUBLESHOOTING:\n" + "\n".join(f"- {tb}" for tb in troubleshooting),
        "KEY LEARNING POINTS:\n" + "\n".join(f"- {l}" for l in learnings)
    ]
    content = "\n\n".join(content_lines)

    return DocumentChunk(
        chunk_id=f"{doc_id}-P{source_metadata.page:02d}-C01",
        document_id=doc_id,
        document_type=DocumentType.OPL,
        title=f"One Point Lesson - {title} ({doc_id})",
        equipment_tag=tag,
        content=content,
        source=source_metadata,
        metadata={
            "discipline": disc,
            "section": "Complete Lesson",
            "related_interlock": interlock_ref,
            "pid_reference": pid_ref
        }
    )


def normalize_opl_excel(file_path: Path) -> List[DocumentChunk]:
    """
    Parses an OPL workbook containing multiple OPL sheets.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    chunks = []

    for idx, sname in enumerate(wb.sheetnames, start=1):
        ws = wb[sname]
        meta: Dict[str, str] = {}
        for r in range(1, 20):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

        doc_id = meta.get("opl_no") or f"OPL-{idx:02d}"
        title = meta.get("opl_title") or sname
        tag = meta.get("equipment_tag")
        disc = meta.get("discipline", "Operations")
        source_file = meta.get("source_file") or f"{doc_id}.pdf"
        # Never infer revision if missing
        rev = normalize_revision(meta.get("revision"))
        status = normalize_status(meta.get("revision_status")) or DocumentStatus.APPROVED

        source = DocumentSource(
            file_name=source_file,
            page=1,
            sheet=sname,
            revision=rev,
            status=status
        )

        purpose = ""
        precautions = []
        tools = []
        steps = []
        troubleshooting = []
        learnings = []

        current_prob: Dict[str, str] = {}
        for r in range(21, ws.max_row + 1):
            sec = str(ws.cell(r, 1).value or "").strip()
            sub = str(ws.cell(r, 2).value or "").strip()
            val = str(ws.cell(r, 3).value or "").strip()

            if "Purpose" in sec:
                purpose = val or sub
            elif "Safety" in sec and val:
                precautions.append(f"{sub}: {val}" if sub else val)
            elif "Tools" in sec and val:
                tools.append(f"{sub}: {val}" if sub else val)
            elif "Procedure" in sec and val:
                steps.append(f"{sub}: {val}")
            elif "Common Problems" in sec:
                if "Symptom" in sub:
                    if current_prob:
                        troubleshooting.append(f"Symptom: {current_prob.get('sym')} | Cause: {current_prob.get('cause')} | Action: {current_prob.get('act')}")
                    current_prob = {"sym": val, "cause": "", "act": ""}
                elif "Cause" in sub:
                    current_prob["cause"] = val
                elif "Action" in sub:
                    current_prob["act"] = val
            elif "Learning" in sec and val:
                learnings.append(f"{sub}: {val}" if sub else val)

        if current_prob:
            troubleshooting.append(f"Symptom: {current_prob.get('sym')} | Cause: {current_prob.get('cause')} | Action: {current_prob.get('act')}")

        sheet_payload = {
            "document_id": doc_id,
            "title": title,
            "equipment_tag": tag,
            "discipline": disc,
            "area_unit": meta.get("area_unit", ""),
            "related_interlock": meta.get("related_interlock", ""),
            "pid_reference": meta.get("pid_reference", ""),
            "sections": {
                "purpose": purpose,
                "safety_precautions": precautions,
                "tools": tools,
                "steps": steps,
                "troubleshooting": troubleshooting,
                "learnings": learnings
            }
        }
        chunks.append(normalize_opl_sheet_dict(sheet_payload, source))

    return chunks
