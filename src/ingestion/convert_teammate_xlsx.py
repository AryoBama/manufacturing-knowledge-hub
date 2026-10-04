import argparse
import datetime
import json
import re
import shutil
import sys
from pathlib import Path
from typing import List, Dict, Any, Optional

import openpyxl

# Ensure project root is in sys.path
BASE_DIR = Path(__file__).resolve().parent.parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from schemas.common import DocumentSource, DocumentStatus, DocumentType
from schemas.document import DocumentChunk
from schemas.maintenance import MaintenanceRecord
from src.ingestion.maintenance_loader import parse_maintenance_excel


def convert_datasheet(file_path: Path) -> DocumentChunk:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    meta: Dict[str, str] = {}
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

    tag = meta.get("equipment_tag", "GA-1201A")
    doc_id = meta.get("doc_no") or meta.get("document_id") or f"TJC-LLD-DS-{tag}"
    rev = meta.get("revision", "Rev 3" if "GA-1201A" in tag else "Rev 0")
    status_str = meta.get("revision_status", "Issued for Operation")
    source_file = meta.get("source_file") or file_path.name.replace(".xlsx", ".pdf")

    # Group parameters by sheet
    sections: Dict[str, List[str]] = {}
    for sname in wb.sheetnames:
        if sname == "Metadata":
            continue
        ws = wb[sname]
        lines = []
        for r in range(2, ws.max_row + 1):
            param = ws.cell(r, 2).value
            val = ws.cell(r, 3).value
            unit = ws.cell(r, 4).value
            if not param:
                param = ws.cell(r, 1).value
                val = ws.cell(r, 2).value
                unit = ws.cell(r, 3).value
            if param and val is not None and str(param).strip().lower() != "parameter":
                unit_str = f" {unit}" if unit and str(unit).strip() not in ("-", "—", "") else ""
                lines.append(f"{str(param).strip()}: {val}{unit_str}")
        if lines:
            sections[sname] = lines

    content_parts = [
        f"EQUIPMENT DATASHEET: {tag} ({meta.get('equipment_name', 'Equipment')}).",
        f"Document No: {doc_id}, Revision: {rev}, Status: {status_str}.",
        f"Equipment Type: {meta.get('equipment_type', 'N/A')}.",
        f"Plant Unit: {meta.get('plant_unit', 'N/A')}, Area: {meta.get('area', 'N/A')}.",
        f"Functional Location: {meta.get('functional_location', 'N/A')}, Criticality: {meta.get('criticality', 'N/A')}.",
        f"Service: {meta.get('service', 'N/A')}."
    ]

    for sname, s_lines in sections.items():
        # Keep canonical header for Mechanical / Motor data
        hdr = "MECHANICAL SPECIFICATIONS:" if "mechanical" in sname.lower() else (
            "MOTOR & ELECTRICAL SPECIFICATIONS:" if ("motor" in sname.lower() or "driver" in sname.lower()) else f"{sname.upper()}:"
        )
        content_parts.append(hdr)
        content_parts.append("; ".join(s_lines))

    pid_ref = meta.get("pid_reference", f"TJC-LLD-PID-{tag.replace('GA-', '')}")
    il_ref = meta.get("interlock_reference", f"SEQ-{tag.split('-')[-1].rstrip('AB')}")
    content_parts.append(f"CROSS REFERENCES: P&ID: {pid_ref}, Interlock: {il_ref}.")
    content = "\n".join(content_parts)

    return DocumentChunk(
        chunk_id=f"{doc_id}-P01-C01",
        document_id=doc_id,
        document_type=DocumentType.DATASHEET,
        title=f"Equipment Datasheet - {meta.get('equipment_name', 'Equipment')} ({tag})",
        equipment_tag=tag,
        content=content,
        source=DocumentSource(
            file_name=source_file,
            page=1,
            revision=rev,
            status=DocumentStatus.ISSUED_FOR_OPERATION
        )
    )


def convert_interlock(file_path: Path) -> DocumentChunk:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    meta: Dict[str, str] = {}
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

    tag = meta.get("equipment_tag", "GA-1201A")
    doc_id = meta.get("doc_no") or meta.get("document_id") or f"TJC-LLD-IL-{tag}"
    rev = meta.get("revision", "Rev 3" if "GA-1201A" in tag else "Rev 0")
    logic_no = meta.get("logic_no", f"SEQ-{tag.split('-')[-1].rstrip('AB')}")
    source_file = meta.get("source_file") or file_path.name.replace(".xlsx", ".pdf")

    trips = []
    if "Cause Effect" in wb.sheetnames:
        ws = wb["Cause Effect"]
        effect_cols = {}
        for c in range(7, ws.max_column + 1):
            col_hdr = ws.cell(2, c).value
            if col_hdr and str(col_hdr).strip():
                raw_name = str(col_hdr).strip()
                eff_match = re.match(r"eff_?(\d+)[_ ]*(.*)", raw_name, re.IGNORECASE)
                if eff_match:
                    num = eff_match.group(1)
                    act_name = eff_match.group(2).replace("_", " ").title()
                    if "trip" in raw_name.lower() and "motor" in raw_name.lower():
                        eff_label = f"Trip Motor (EFF-{num})"
                    elif "close" in raw_name.lower() and "xv" in raw_name.lower():
                        eff_label = f"Close Discharge XV-1201 (EFF-{num})"
                    elif "open" in raw_name.lower() and "fv" in raw_name.lower():
                        eff_label = f"Open Min-Flow FV-1201 (EFF-{num})"
                    elif "alarm" in raw_name.lower():
                        eff_label = f"DCS Alarm (EFF-{num})"
                    elif "standby" in raw_name.lower():
                        eff_label = f"Start Standby Pump GA-1201B (EFF-{num})"
                    else:
                        eff_label = f"{act_name} (EFF-{num})"
                else:
                    eff_label = raw_name.replace("_", " ").title()
                effect_cols[c] = eff_label

        for r in range(3, ws.max_row + 1):
            trip_id = ws.cell(r, 2).value
            desc = ws.cell(r, 3).value
            tag_inst = ws.cell(r, 4).value
            sp = ws.cell(r, 5).value
            voting = ws.cell(r, 6).value
            if trip_id and str(trip_id).strip().startswith("T") and len(str(trip_id).strip()) <= 4:
                effs = []
                for c, eff_label in effect_cols.items():
                    val = str(ws.cell(r, c).value or "").strip().upper()
                    if val == "X":
                        effs.append(eff_label)
                actions = ", ".join(effs)
                trips.append(f"- {trip_id}: {desc} ({tag_inst} {sp}, {voting}) -> Actions: {actions}")

    permissives = []
    if "Start Permissive" in wb.sheetnames:
        ws = wb["Start Permissive"]
        for r in range(3, ws.max_row + 1):
            p_id = ws.cell(r, 2).value
            desc = ws.cell(r, 3).value
            sig = ws.cell(r, 4).value
            gate = ws.cell(r, 5).value
            if p_id and desc:
                permissives.append(f"- {p_id}: {desc} ({sig}, Gate Logic: {gate})")

    pid_ref = meta.get("pid_reference", f"TJC-LLD-PID-{tag.replace('GA-', '')}")
    ds_ref = meta.get("datasheet_reference", f"TJC-LLD-DS-{tag}")

    content_parts = [
        f"INTERLOCK LOGIC & CAUSE EFFECT MATRIX: {logic_no} ({meta.get('logic_description', 'Shutdown Logic')}).",
        f"Document No: {doc_id}, Revision: {rev}, SIL Level: {meta.get('sil_level', 'SIL 1')}.",
        f"Equipment: {tag}, Functional Location: {meta.get('functional_location', 'N/A')}.",
        "CAUSE & EFFECT MATRIX (TRIPS):",
        "\n".join(trips),
        "START PERMISSIVES (AND-gate conditions to start pump):",
        "\n".join(permissives),
        f"CROSS REFERENCES: P&ID: {pid_ref}, Datasheet: {ds_ref}."
    ]
    content = "\n".join(content_parts)

    return DocumentChunk(
        chunk_id=f"{doc_id}-P01-C01",
        document_id=doc_id,
        document_type=DocumentType.INTERLOCK,
        title=f"Interlock Logic & Cause Effect Matrix - {tag} ({logic_no})",
        equipment_tag=tag,
        content=content,
        source=DocumentSource(
            file_name=source_file,
            page=1,
            revision=rev,
            status=DocumentStatus.ISSUED_FOR_OPERATION
        )
    )


def convert_pid(file_path: Path) -> DocumentChunk:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    meta: Dict[str, str] = {}
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

    tag = meta.get("equipment_tag", "GA-1201A")
    raw_doc_id = meta.get("doc_no") or meta.get("document_id")
    if raw_doc_id and "XXXX" not in raw_doc_id:
        doc_id = raw_doc_id.rstrip("AB") if (raw_doc_id.endswith("A") or raw_doc_id.endswith("B")) else raw_doc_id
    else:
        doc_id = f"TJC-LLD-PID-{tag.replace('GA-', '').rstrip('AB')}"
    rev = meta.get("revision", "Rev 0")
    source_file = meta.get("source_file") or file_path.name.replace(".xlsx", ".png")

    inst_lines = []
    if "Instrument List" in wb.sheetnames:
        ws = wb["Instrument List"]
        for r in range(3, ws.max_row + 1):
            itag = ws.cell(r, 1).value
            desc = ws.cell(r, 2).value
            itype = ws.cell(r, 3).value
            sp = ws.cell(r, 4).value
            unit = ws.cell(r, 5).value
            fn = ws.cell(r, 6).value
            rel_il = ws.cell(r, 7).value
            loc = ws.cell(r, 8).value
            if itag:
                sp_str = f" [Setpoint: {sp} {unit}]" if sp and str(sp).strip() not in ("-", "—", "") else ""
                inst_lines.append(f"- {itag}: {desc} ({itype}){sp_str}. Function: {fn}. Location: {loc}. Interlock: {rel_il}")

    ds_ref = meta.get("datasheet_reference", f"TJC-LLD-DS-{tag}")
    il_ref = meta.get("interlock_reference", f"SEQ-{tag.split('-')[-1].rstrip('AB')}")

    content_parts = [
        f"PIPING & INSTRUMENTATION DIAGRAM (P&ID) REGISTER: {tag} ({meta.get('equipment_name', 'Equipment')}).",
        f"Document No: {doc_id}, Title: {meta.get('title', f'{tag} System')}.",
        f"Area: {meta.get('area', '')}, Functional Location: {meta.get('functional_location', '')}.",
        "INSTRUMENT & PROCESS CONTROL LOOPS:",
        "\n".join(inst_lines),
        f"CROSS REFERENCES: Datasheet: {ds_ref}, Interlock: {il_ref}."
    ]
    content = "\n".join(content_parts)

    return DocumentChunk(
        chunk_id=f"{doc_id}-P01-C01",
        document_id=doc_id,
        document_type=DocumentType.PID,
        title=f"Piping & Instrumentation Diagram Register - {tag}",
        equipment_tag=tag,
        content=content,
        source=DocumentSource(
            file_name=source_file,
            page=1,
            revision=rev,
            status=DocumentStatus.ISSUED_FOR_OPERATION
        )
    )


def convert_plot_plan(file_path: Path) -> DocumentChunk:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    meta: Dict[str, str] = {}
    if "Metadata" in wb.sheetnames:
        ws = wb["Metadata"]
        for r in range(1, ws.max_row + 1):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

    tag = meta.get("equipment_tag", "GA-1201A")
    raw_doc_id = meta.get("doc_no") or meta.get("document_id")
    doc_id = raw_doc_id or f"TJC-LLD-PP-{tag}"
    rev = meta.get("revision", "Rev 0")
    source_file = meta.get("source_file") or file_path.name.replace(".xlsx", ".pdf")

    loc_params = []
    nearby = []
    if "Location Data" in wb.sheetnames:
        ws = wb["Location Data"]
        for r in range(3, 15):
            param = ws.cell(r, 1).value
            val = ws.cell(r, 2).value
            unit = ws.cell(r, 3).value
            if param and val:
                unit_str = f" {unit}" if unit and str(unit).strip() not in ("-", "—", "") else ""
                loc_params.append(f"{param}: {val}{unit_str}")
        for r in range(17, ws.max_row + 1):
            etag = ws.cell(r, 1).value
            ename = ws.cell(r, 2).value
            area_code = ws.cell(r, 3).value
            area_name = ws.cell(r, 4).value
            el = ws.cell(r, 5).value
            if etag:
                name_str = f" ({ename})" if ename and str(ename).strip() not in ("-", "—", "") else ""
                nearby.append(f"- {etag}{name_str}: Area {area_code} ({area_name}), Elevation {el}")

    content_parts = [
        f"PLOT PLAN & EQUIPMENT LOCATION DATA: {tag} ({meta.get('equipment_name', 'Equipment')}).",
        f"Document No: {doc_id}, Revision: {rev}, Status: {meta.get('revision_status', 'Issued for Construction')}.",
        f"Area: {meta.get('area', '')}, Functional Location: {meta.get('functional_location', '')}.",
        "LOCATION SPECIFICATIONS & COORDINATES:",
        "; ".join(loc_params),
        "NEARBY PLANT EQUIPMENT & ELEVATIONS:",
        "\n".join(nearby)
    ]
    content = "\n".join(content_parts)

    return DocumentChunk(
        chunk_id=f"{doc_id}-P01-C01",
        document_id=doc_id,
        document_type=DocumentType.PLOT_PLAN,
        title=f"Plot Plan & Location Data - {tag}",
        equipment_tag=tag,
        content=content,
        source=DocumentSource(
            file_name=source_file,
            page=1,
            revision=rev,
            status=DocumentStatus.ISSUED_FOR_CONSTRUCTION
        )
    )


def convert_opl(file_path: Path) -> List[DocumentChunk]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    chunks = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        meta: Dict[str, str] = {}
        for r in range(1, 20):
            k = ws.cell(r, 1).value
            v = ws.cell(r, 2).value
            if k and v:
                meta[str(k).strip()] = str(v).strip()

        doc_id = meta.get("opl_no")
        if not doc_id:
            continue
        title = meta.get("opl_title", sheet_name)
        tag = meta.get("equipment_tag", "GA-1201A")
        disc = meta.get("discipline", "Mechanical")
        source_file = meta.get("source_file", f"{doc_id}.pdf")

        purpose = ""
        precautions = []
        tools = []
        steps = []
        problems = []
        learnings = []

        current_problem = {}

        for r in range(21, ws.max_row + 1):
            sec = ws.cell(r, 1).value
            sub = ws.cell(r, 2).value
            val = ws.cell(r, 3).value
            sec_str = str(sec).strip() if sec else ""
            sub_str = str(sub).strip() if sub else ""
            val_str = str(val).strip() if val else ""

            if "Purpose" in sec_str:
                purpose = val_str or sub_str
            elif "Safety" in sec_str:
                if val_str:
                    precautions.append(f"- {sub_str}: {val_str}" if sub_str else f"- {val_str}")
            elif "Tools" in sec_str:
                if val_str:
                    tools.append(f"- {sub_str}: {val_str}" if sub_str else f"- {val_str}")
            elif "Procedure" in sec_str:
                if val_str:
                    steps.append(f"- {sub_str}: {val_str}" if sub_str else f"- {val_str}")
            elif "Common Problems" in sec_str:
                if "Symptom" in sub_str:
                    if current_problem:
                        problems.append(current_problem)
                    current_problem = {"symptom": val_str, "cause": "", "action": ""}
                elif "Cause" in sub_str:
                    current_problem["cause"] = val_str
                elif "Action" in sub_str:
                    current_problem["action"] = val_str
            elif "Learning" in sec_str:
                if val_str:
                    learnings.append(f"- {sub_str}: {val_str}" if sub_str else f"- {val_str}")

        if current_problem:
            problems.append(current_problem)

        prob_lines = [
            f"- Symptom: {p.get('symptom')} | Likely Cause: {p.get('cause')} | Action: {p.get('action')}"
            for p in problems if p.get("symptom")
        ]

        content_parts = [
            f"ONE POINT LESSON (OPL): {doc_id} - {title}.",
            f"Equipment: {tag} ({meta.get('equipment_name', 'Equipment')}), Discipline: {disc}, Area: {meta.get('area_unit', '')}.",
            f"Related Interlock: {meta.get('related_interlock', '')}, P&ID: {meta.get('pid_reference', '')}.",
            f"PURPOSE & OBJECTIVE:\n{purpose}",
            "SAFETY PRECAUTIONS:\n" + "\n".join(precautions),
            "TOOLS & MATERIALS:\n" + "\n".join(tools),
            "STEP-BY-STEP PROCEDURE:\n" + "\n".join(steps),
            "COMMON PROBLEMS & TROUBLESHOOTING:\n" + "\n".join(prob_lines),
            "KEY LEARNING POINTS:\n" + "\n".join(learnings)
        ]
        content = "\n\n".join(content_parts)

        chunk = DocumentChunk(
            chunk_id=f"{doc_id}-P01-C01",
            document_id=doc_id,
            document_type=DocumentType.OPL,
            title=f"One Point Lesson - {title} ({doc_id})",
            equipment_tag=tag,
            content=content,
            source=DocumentSource(
                file_name=source_file,
                page=1,
                revision="Rev 0",
                status=DocumentStatus.ISSUED_FOR_OPERATION
            )
        )
        chunks.append(chunk)

    return chunks


def _find_file(dir_path: Path, patterns: List[str]) -> Optional[Path]:
    for pat in patterns:
        matches = list(dir_path.glob(pat))
        if matches:
            return matches[0]
    return None


def convert_all(
    input_dir: Path,
    output_dir: Path,
    copy_to_input_folder: bool = False,
    target_tag: Optional[str] = None
) -> Dict[str, Any]:
    print(f"Converting Excel files from: {input_dir}")
    print(f"Target extracted directory: {output_dir}")

    output_dir.mkdir(parents=True, exist_ok=True)
    opl_dir = output_dir / "opl"
    opl_dir.mkdir(parents=True, exist_ok=True)

    results_summary = {}

    # 1. Datasheet
    ds_file = _find_file(input_dir, ["*Datasheet*.xlsx", "*datasheet*.xlsx"])
    detected_tag = target_tag
    if ds_file and ds_file.exists():
        print(f"  [+] Parsing Datasheet: {ds_file.name}")
        ds_chunk = convert_datasheet(ds_file)
        if not detected_tag and ds_chunk.equipment_tag:
            detected_tag = ds_chunk.equipment_tag
        out_ds = output_dir / "datasheet.json"
        with open(out_ds, "w", encoding="utf-8") as f:
            json.dump(ds_chunk.model_dump(mode="json"), f, indent=2)
        results_summary["datasheet"] = str(out_ds)

    # 2. Interlock
    il_file = _find_file(input_dir, ["*Interlock*.xlsx", "*interlock*.xlsx"])
    if il_file and il_file.exists():
        print(f"  [+] Parsing Interlock: {il_file.name}")
        il_chunk = convert_interlock(il_file)
        out_il = output_dir / "interlock.json"
        with open(out_il, "w", encoding="utf-8") as f:
            json.dump(il_chunk.model_dump(mode="json"), f, indent=2)
        results_summary["interlock"] = str(out_il)

    # 3. PID Register
    pid_file = _find_file(input_dir, ["*PID*.xlsx", "*P&ID*.xlsx", "*pid*.xlsx"])
    if pid_file and pid_file.exists():
        print(f"  [+] Parsing PID Register: {pid_file.name}")
        pid_chunk = convert_pid(pid_file)
        out_pid = output_dir / "pid.json"
        with open(out_pid, "w", encoding="utf-8") as f:
            json.dump(pid_chunk.model_dump(mode="json"), f, indent=2)
        results_summary["pid"] = str(out_pid)

    # 4. Plot Plan
    pp_file = _find_file(input_dir, ["*Plot Plan*.xlsx", "*plot plan*.xlsx", "*Plot*.xlsx"])
    if pp_file and pp_file.exists():
        print(f"  [+] Parsing Plot Plan: {pp_file.name}")
        pp_chunk = convert_plot_plan(pp_file)
        out_pp = output_dir / "plot_plan.json"
        with open(out_pp, "w", encoding="utf-8") as f:
            json.dump(pp_chunk.model_dump(mode="json"), f, indent=2)
        results_summary["plot_plan"] = str(out_pp)

    # 5. OPL (7 sheets)
    opl_file = _find_file(input_dir, ["*OPL*.xlsx", "*opl*.xlsx"])
    if opl_file and opl_file.exists():
        print(f"  [+] Parsing OPL Sheets: {opl_file.name}")
        opl_chunks = convert_opl(opl_file)
        saved_opls = []
        for c in opl_chunks:
            chunk_file = opl_dir / f"{c.document_id}.json"
            with open(chunk_file, "w", encoding="utf-8") as f:
                json.dump(c.model_dump(mode="json"), f, indent=2)
            saved_opls.append(str(chunk_file))
        results_summary["opl_chunks"] = len(opl_chunks)

    # 6. Maintenance History
    mnt_file = _find_file(input_dir, ["*Maintenance History*.xlsx", "*maintenance*.xlsx"])
    if mnt_file and mnt_file.exists():
        print(f"  [+] Parsing Maintenance History: {mnt_file.name}")
        records = parse_maintenance_excel(mnt_file)
        # Filter by detected tag if provided to avoid 211 events replicated across all tags
        if detected_tag:
            tag_records = [r for r in records if r.equipment_tag.upper() == detected_tag.upper()]
            if tag_records:
                records = tag_records
        out_mnt = output_dir / "maintenance_history.json"
        with open(out_mnt, "w", encoding="utf-8") as f:
            json.dump([r.model_dump(mode="json") for r in records], f, indent=2)
        results_summary["maintenance_records"] = len(records)

    # Optionally copy to input folder under json/
    if copy_to_input_folder:
        input_json_dir = input_dir / "json"
        print(f"  [+] Mirroring JSON copies to: {input_json_dir}")
        input_json_dir.mkdir(parents=True, exist_ok=True)
        for item in output_dir.iterdir():
            dest = input_json_dir / item.name
            if item.is_dir():
                if dest.exists():
                    shutil.rmtree(dest)
                shutil.copytree(item, dest)
            else:
                shutil.copy2(item, dest)

    print("\nConversion completed successfully!")
    return results_summary


def convert_all_equipment(
    raw_dir: Path,
    extracted_base_dir: Path
) -> Dict[str, Any]:
    """
    Converts Excel files for all 8 equipment packages found in raw_dir.
    """
    equipment_map = {
        "GA-1201A": "GA-1201A",
        "KC-4501": "KC-4501",
        "YD-2301": "YD-2301",
        "DC-3401A": "DC-3401A",
        "EA-5601": "EA-5601",
        "LV-6701": "LV-6701",
        "CT-7801": "CT-7801",
        "FA-8901": "FA-8901"
    }

    all_summaries = {}
    for d in sorted(raw_dir.iterdir()):
        if not d.is_dir():
            continue
        tag = None
        for k in equipment_map.keys():
            if k in d.name.upper():
                tag = k
                break
        if not tag:
            continue

        target_extracted = extracted_base_dir / tag
        print(f"\n=======================================================")
        print(f"PROCESSING RAW PACKAGE: {d.name} -> TAG: {tag}")
        print(f"=======================================================")
        summary = convert_all(input_dir=d, output_dir=target_extracted, target_tag=tag)
        all_summaries[tag] = summary

        # For GA-1201A, also mirror to Set_01_GA-1201A for backwards compatibility
        if tag == "GA-1201A":
            legacy_dir = extracted_base_dir / "Set_01_GA-1201A"
            legacy_dir.mkdir(parents=True, exist_ok=True)
            for item in target_extracted.iterdir():
                dest = legacy_dir / item.name
                if item.is_dir():
                    if dest.exists():
                        shutil.rmtree(dest)
                    shutil.copytree(item, dest)
                else:
                    shutil.copy2(item, dest)
            print(f"  [+] Mirrored GA-1201A to legacy folder: {legacy_dir}")

    return all_summaries


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert teammate Excel files to AI contract JSON")
    parser.add_argument("--input-dir", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--all", action="store_true", help="Convert all equipment in raw folder")
    args = parser.parse_args()

    raw_base = BASE_DIR / "data" / "raw"
    ext_base = BASE_DIR / "data" / "extracted"

    if args.all or (args.input_dir is None and args.output_dir is None):
        if raw_base.exists():
            convert_all_equipment(raw_base, ext_base)
        else:
            convert_all_equipment(BASE_DIR, ext_base)
    else:
        inp = args.input_dir or (BASE_DIR / "data" / "raw" / "GA-1201A")
        outp = args.output_dir or (ext_base / "GA-1201A")
        convert_all(inp, outp)
